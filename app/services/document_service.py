import hashlib
import os
import shutil
from contextlib import contextmanager
from pathlib import Path
from uuid import uuid4
from zipfile import ZipFile

from docx import Document
from fastapi import HTTPException
from pypdf import PdfReader
from sqlalchemy.orm import Session

from app.models.support import KnowledgeDocument, DocumentChunk
from app.rag.ollama import ProviderError, embed, embedding_model
from app.rag.vector_store import document_lock, vector_store

ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".markdown"}
MAX_FILE_SIZE = 10 * 1024 * 1024
CHUNK_SIZE = 180
CHUNK_OVERLAP = 30


def knowledge_root():
    return Path(os.getenv("KNOWLEDGE_PATH", "data/knowledge_base"))


@contextmanager
def document_mutation():
    if not document_lock.acquire(blocking=False):
        raise HTTPException(409, "Kho tri thức đang xử lý tài liệu. Vui lòng thử lại sau.")
    try:
        yield
    finally:
        document_lock.release()


def chunk_text(text):
    words = text.split()
    chunks = []
    for start in range(0, len(words), CHUNK_SIZE - CHUNK_OVERLAP):
        chunks.append(" ".join(words[start:start + CHUNK_SIZE]))
        if start + CHUNK_SIZE >= len(words):
            break
    return chunks


def extract_sections(path):
    if path.suffix.lower() == ".pdf":
        reader = PdfReader(str(path))
        if reader.is_encrypted or len(reader.pages) > 1000:
            raise ValueError("PDF được mã hóa hoặc vượt 1.000 trang.")
        for number, page in enumerate(reader.pages, 1):
            yield page.extract_text() or "", number, f"Trang {number}"
    elif path.suffix.lower() == ".docx":
        with ZipFile(path) as archive:
            if sum(item.file_size for item in archive.infolist()) > 50 * 1024 * 1024:
                raise ValueError("DOCX giải nén vượt 50 MB.")
        document = Document(str(path))
        for number, paragraph in enumerate(document.paragraphs, 1):
            yield paragraph.text, None, f"Đoạn {number}"
        for number, table in enumerate(document.tables, 1):
            for row_number, row in enumerate(table.rows, 1):
                yield " | ".join(cell.text for cell in row.cells), None, f"Bảng {number}, hàng {row_number}"
    else:
        for number, paragraph in enumerate(path.read_text(encoding="utf-8-sig").splitlines(), 1):
            yield paragraph, None, f"Dòng {number}"


def document_info(db, record):
    return {"document_id": record.id, "filename": record.filename, "status": record.status,
            "error_message": record.error_message, "embedding_model": record.embedding_model,
            "created_at": record.created_at,
            "chunks": db.query(DocumentChunk).filter_by(document_id=record.id).count()}


def index_document(db, record, root=None):
    document_id, filename = record.id, record.filename
    root = root or knowledge_root()
    record.status, record.error_message, record.index_version = "processing", None, None
    db.commit()
    try:
        chunks = []
        version = str(uuid4())
        characters = 0
        for text, page, location in extract_sections(root / document_id / filename):
            characters += len(text)
            if characters > 2_000_000:
                raise ValueError("Tài liệu vượt 2 triệu ký tự trích xuất.")
            for content in chunk_text(text):
                chunks.append({"chunk_id": f"{document_id}_{len(chunks):04d}", "document_id": document_id,
                               "index_version": version, "content": content, "source": filename,
                               "page": page, "location": location})
        if not chunks:
            raise ValueError("Không trích xuất được chữ. PDF scan cần OCR trước khi tải lên.")
        if len(chunks) > 3000:
            raise ValueError("Tài liệu vượt 3.000 đoạn; hãy chia nhỏ file.")
        db.query(DocumentChunk).filter_by(document_id=document_id).delete()
        for index, chunk in enumerate(chunks):
            db.add(DocumentChunk(id=chunk["chunk_id"], document_id=document_id, chunk_index=index,
                                 content=chunk["content"], source=filename, page=chunk["page"], location=chunk["location"]))
        db.commit()
        vector_store.delete(document_id)
        for start in range(0, len(chunks), 16):
            batch = chunks[start:start + 16]
            vector_store.upsert(batch, embed([f"title: {filename} | text: {c['content']}" for c in batch]))
        record = db.get(KnowledgeDocument, document_id)
        record.status, record.index_version, record.embedding_model = "indexed", version, embedding_model()
        db.commit()
    except Exception as exc:
        db.rollback()
        record = db.get(KnowledgeDocument, document_id)
        record.status = "failed"
        record.error_message = str(exc) if isinstance(exc, (ValueError, ProviderError)) else "Không xử lý được tài liệu. Kiểm tra file và kho vector rồi thử lập chỉ mục lại."
        db.commit()
    return document_info(db, record)


def save_document(db: Session, filename: str, content: bytes, root: Path) -> dict:
    filename = Path(filename.replace("\\", "/")).name
    if (Path(filename).suffix.lower() not in ALLOWED_EXTENSIONS or len(filename) > 255 or
            any(ord(c) < 32 or c in '<>:"|?*' for c in filename) or Path(filename).is_reserved()):
        raise ValueError("Chỉ hỗ trợ PDF, DOCX, TXT và Markdown; tên file tối đa 255 ký tự.")
    if not content or len(content) > MAX_FILE_SIZE:
        raise ValueError("File phải có nội dung và không vượt 10 MB.")
    fingerprint = hashlib.sha256(content).hexdigest()
    with document_mutation():
        if db.query(KnowledgeDocument).filter_by(file_hash=fingerprint).first():
            raise HTTPException(409, "Nội dung tài liệu đã tồn tại. Dùng lập chỉ mục lại nếu cần.")
        document_id = str(uuid4())
        target_dir = root / document_id
        target_dir.mkdir(parents=True)
        try:
            (target_dir / filename).write_bytes(content)
            record = KnowledgeDocument(id=document_id, filename=filename, file_hash=fingerprint, status="processing")
            db.add(record)
            db.commit()
        except Exception:
            db.rollback()
            shutil.rmtree(target_dir)
            raise
        return index_document(db, record, root)


def delete_document(db, record):
    document_id = record.id
    # Hide first: interrupted cleanup must never leave searchable deleted content.
    record.status, record.index_version = "deleting", None
    db.commit()
    try:
        vector_store.delete(document_id)
        directory = knowledge_root() / document_id
        if directory.exists():
            shutil.rmtree(directory)
        db.query(DocumentChunk).filter_by(document_id=document_id).delete()
        db.delete(record)
        db.commit()
    except Exception as exc:
        db.rollback()
        raise HTTPException(503, "Xóa chưa hoàn tất. Tài liệu đã ngừng được truy xuất; hãy thử xóa lại.") from exc
