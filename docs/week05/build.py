"""Build Week 05 UML figures and report. Run --check for a read-only self-check.

ponytail: fixed layouts for this report; use PlantUML editor for larger models.
"""
import argparse
from html import escape
import importlib.util
import math
from pathlib import Path
import re
from xml.etree import ElementTree as ET
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[2]
HERE = Path(__file__).resolve().parent
OUT = HERE / 'diagrams'
UC = {
    'UC01':'Cấp tài khoản nhân viên', 'UC02':'Đăng nhập và kết thúc phiên nhân viên',
    'UC03':'Sử dụng hội thoại Website riêng', 'UC04':'Hỏi chính sách có nguồn',
    'UC05':'Quản lý kho tri thức', 'UC06':'Yêu cầu hỗ trợ từ nhân viên',
    'UC07':'Tiếp nhận yêu cầu đang chờ', 'UC08':'Phản hồi khách đang phụ trách',
    'UC09':'Hoàn tất lượt hỗ trợ', 'UC10':'Tiếp tục hỗ trợ sau giải quyết',
    'UC11':'Theo dõi Inbox và SLA', 'UC12':'Duyệt chất lượng câu trả lời',
    'UC13':'Tra cứu đơn được phép', 'UC14':'Xử lý phản hồi Telegram gửi lỗi',
    'UC15':'Quản lý tài khoản khách', 'UC16':'Xem hồ sơ đơn và báo cáo',
    'UC17':'Sao lưu và khôi phục dữ liệu', 'UC18':'Đọc nguồn và hỏi thử nội bộ',
    'UC19':'Cấp hoặc thu hồi quyền tra đơn', 'UC20':'Quản lý dữ liệu bán lẻ giả lập',
}
ACTORS = {'KH':'Khách hàng','NV':'Nhân viên','AD':'Quản trị viên','TG':'Telegram Bot API',
          'SMTP':'SMTP','RV':'Người duyệt','OP':'Người vận hành'}
ASSOCIATIONS = {
    'KH':['UC03','UC04','UC06','UC10','UC13','UC15'],
    'NV':['UC02','UC06','UC07','UC08','UC09','UC11','UC14','UC16','UC18'],
    'AD':['UC01','UC05','UC19','UC20'], 'TG':['UC04','UC06','UC08','UC10','UC13','UC14'],
    'SMTP':['UC15'], 'RV':['UC12'], 'OP':['UC17'],
}
MODULES = {
    'M01':('Unified Inbox và ticket',['UC07','UC08','UC09','UC10','UC11','UC14']),
    'M02':('Kho tri thức',['UC05','UC18']),
    'M03':('Hỏi đáp và tra đơn',['UC04','UC13']),
    'M04':('Chuyển giao nhân viên',['UC06']),
    'S01':('Truy cập và danh tính',['UC01','UC02','UC03','UC15','UC19']),
    'S02':('Quản lý đánh giá và vận hành',['UC12','UC16','UC17','UC20']),
}
# Events: message/return, note, fragment open/else/end. Aliases refer to real responsibilities.
SEQUENCES = {
 'SD01':('Hỏi chính sách và kiểm tra AI lỗi thời',[
  ('UI','Widget','Client'),('C','WidgetController','Controller'),('M','MessageService','Service'),
  ('R','RAGService','Service'),('DB','SQLite','Database'),('O','Ollama','Local API'),('Q','Qdrant','Vector store')],[
  ('msg','UI','C','Gửi câu hỏi và UUID; phiên/CSRF/rate limit'),
  ('msg','C','M','process_message sau khi có slot AI'),
  ('msg','M','DB','Khóa ghi; kiểm tra UUID; lưu tin; COMMIT'),
  ('msg','M','R','answer_question với lịch sử đúng hội thoại'),
  ('msg','R','DB','Đọc phiên bản indexed; COMMIT'),
  ('msg','R','O','Embedding câu hỏi'),('msg','R','Q','Truy hồi theo phiên bản được phép'),
  ('msg','R','O','Chọn ID bằng chứng và kiểm định có giới hạn'),
  ('open','alt','[ProviderError]'),('ret','R','M','Lỗi provider; tin khách đã lưu'),
  ('ret','M','C','ai_error; giữ tin khách'),
  ('else','[Có đáp án hoặc từ chối có kiểm soát]'),('ret','R','M','Ứng viên, citations, reviewed_sources'),
  ('open','critical','[Kiểm tra lại sau model]'),
  ('note','document_lock khi có đáp án; transaction ghi SQLite ngắn'),
  ('msg','M','DB','Đọc lại status, tin cuối và nguồn hiện tại'),
  ('open','alt','[open; đúng tin cuối; nguồn hợp lệ]'),('msg','M','DB','Lưu AI + reply_to_id; COMMIT'),
  ('else','[Handoff, tin mới hoặc nguồn đổi]'),('msg','M','DB','Bỏ AI cũ; COMMIT, giữ lịch sử'),('end',),('end',),
  ('ret','M','C','Kết quả đã lọc thông tin nội bộ'),('end',),('ret','C','UI','Snapshot hoặc lỗi; giải phóng slot'),
 ]),
 'SD02':('Cấp quyền và tra đơn chỉ đọc',[
  ('A','Quản lý Admin','Client'),('W','Website khách','Client'),('C','Workspace/Widget','Controller'),
  ('M','MessageService','Service'),('T','ToolService','Service'),('S','OrderService','Service'),
  ('DB','SQLite','Database'),('O','Ollama','Local API')],[
  ('note','Admin xác minh qua kênh tin cậy ngoài ứng dụng; mã bí mật không đi vào model'),
  ('msg','A','C','Cấp access-code sau xác minh; quyền admin'),
  ('msg','C','DB','Khóa Order; xoay mã hash/hạn 15 phút; COMMIT'),
  ('ret','C','A','Mã chỉ hiện cho admin để giao kênh tin cậy'),
  ('msg','W','C','Nhập mã qua biểu mẫu riêng'),
  ('msg','C','DB','Kiểm tra hash/chủ/hạn/phiên; ràng buộc; COMMIT'),
  ('msg','W','C','Hỏi trạng thái một đơn; UUID'),('msg','C','M','process_message theo danh tính server'),
  ('msg','M','DB','Lưu tin, mã từ tin/lịch sử khách; COMMIT'),
  ('msg','M','T','select_order_tool với mã server chọn'),('msg','T','O','JSON Schema: lookup_order/handoff/rag'),
  ('ret','T','M','Lựa chọn đã kiểm tra schema'),
  ('open','critical','[Nhánh lookup_order còn là tin hiện tại]'),
  ('msg','M','DB','Khóa hội thoại; kiểm tra status và tin cuối'),('msg','M','S','lookup_order theo customer_id server'),
  ('msg','S','DB','Kiểm tra chủ hoặc quyền hiện tại; đọc đơn'),
  ('open','alt','[Có quyền và tìm thấy]'),('ret','S','M','Chỉ status và tracking_code'),
  ('msg','M','DB','Lưu phản hồi chỉ đọc; COMMIT'),
  ('else','[Không đủ quyền hoặc không tìm thấy]'),('msg','M','DB','Queue handoff, không lộ đơn; COMMIT'),('end',),('end',),
  ('ret','M','C','Phản hồi hoặc hàng chờ'),('ret','C','W','Snapshot công khai'),
 ]),
 'SD03':('Tiếp nhận phản hồi và hoàn tất ticket',[
  ('UI','Inbox nhân viên','Client'),('C','InboxController','Controller'),('T','Ticket/SLA','Service'),
  ('DB','SQLite','Database'),('W','WidgetController','Controller'),('KH','Website khách','Client')],[
  ('msg','UI','C','Tiếp nhận; phiên staff và CSRF'),
  ('msg','C','DB','UPDATE chờ + chưa phụ trách -> assigned'),
  ('open','alt','[rowcount = 1]'),('msg','C','DB','Đổi ticket open -> assigned; COMMIT'),
  ('ret','C','UI','Tiếp nhận thành công'),
  ('else','[Người khác đã nhận]'),('msg','C','DB','ROLLBACK'),('ret','C','UI','409; đọc lại trước thao tác khác'),('end',),
  ('open','opt','[Đã là người phụ trách]'),
  ('msg','UI','C','Gửi phản hồi với UUID ổn định'),
  ('msg','C','DB','Kiểm tra phụ trách; tái dùng/lưu tin; COMMIT'),
  ('msg','KH','W','Polling lấy lịch sử phiên'),('msg','W','DB','Đọc đúng hội thoại; lọc dữ liệu công khai'),
  ('ret','W','KH','Phản hồi đã lưu'),
  ('msg','UI','C','Finish: ticket, tin cuối, status và ghi chú'),
  ('open','critical','[Transaction hoàn tất]'),('msg','C','DB','Khóa ghi; đọc ticket/người/tin cuối'),
  ('open','alt','[Đúng lượt, phụ trách và tin cuối]'),('msg','C','T','complete_tickets và tính snapshot SLA'),
  ('msg','T','DB','Ghi kết quả, người, thời điểm và SLA'),
  ('msg','C','DB','Đổi status + tin hệ thống; COMMIT'),('ret','C','UI','Hoàn tất đúng lượt'),
  ('else','[Tin mới hoặc điều kiện không còn đúng]'),('msg','C','DB','ROLLBACK'),('ret','C','UI','409/404; đọc lại'),
  ('end',),('end',),('end',),
 ]),
 'SD04':('Telegram nhận giao và phục hồi tin',[
  ('UI','Inbox nhân viên','Client'),('C','InboxController','Controller'),('W','TelegramWorker','Adapter'),
  ('M','MessageService','Service'),('DB','SQLite','Database'),('TG','Telegram Bot API','Third party')],[
  ('msg','W','TG','getUpdates(offset đã commit)'),
  ('msg','W','DB','Lưu update theo bot:ID và cursor; COMMIT'),
  ('msg','W','M','process_next -> process_message, cùng ID'),
  ('note','RAG dùng SD01; /human đi handoff. Update lặp không nhân bản tin/ticket.'),
  ('msg','W','DB','Update done; queue_outgoing pending; COMMIT'),
  ('open','loop','[Tin đầu chưa hoàn tất của từng peer]'),
  ('msg','W','DB','CAS pending -> sending; kiểm tra AI; COMMIT'),
  ('msg','W','TG','sendMessage; không giữ SQL transaction'),
  ('open','alt','[Xác nhận message_id hợp lệ]'),('msg','W','DB','sent + external ID + sent_at; COMMIT'),
  ('else','[Lỗi xác định hoặc kết quả không chắc]'),('msg','W','DB','failed/uncertain; handoff khi open; COMMIT'),
  ('note','failed/uncertain chặn tin sau cùng peer. AI stale -> skipped trước gọi API.'),('end',),('end',),
  ('open','opt','[Khởi động lại sau sending dở]'),('msg','W','DB','recover_sends: sending -> uncertain; COMMIT'),('end',),
  ('open','opt','[Nhân viên quyết định retry hoặc skip]'),
  ('msg','UI','C','retry-delivery; xác nhận uncertain nếu retry'),
  ('msg','C','DB','Kiểm tra phụ trách/trạng thái; pending/skipped; COMMIT'),
  ('ret','C','UI','Trạng thái mới; retry chưa có nghĩa đã giao'),('end',),
 ]),
 'SD05':('Lập chỉ mục nguồn và xử lý lỗi',[
  ('UI','Kho tri thức','Client'),('C','DocumentsController','Controller'),('S','DocumentService','Service'),
  ('F','File tài liệu','File store'),('DB','SQLite','Database'),('O','Ollama','Local API'),('Q','Qdrant','Vector store')],[
  ('msg','UI','C','Upload; admin/CSRF; tối đa 10 MB'),('msg','C','S','save_document(file, nội dung)'),
  ('open','critical','[document_lock trong một tiến trình]'),
  ('msg','S','DB','Kiểm tra SHA-256 chưa tồn tại'),('msg','S','F','Lưu file tài liệu'),
  ('msg','S','DB','processing, index_version=null; COMMIT'),
  ('open','alt','[Đọc và lập chỉ mục thành công]'),
  ('msg','S','F','Đọc chữ và chia đoạn có vị trí nguồn'),
  ('msg','S','DB','Lưu các đoạn nguồn; COMMIT'),('msg','S','Q','Dọn vector cũ của tài liệu'),
  ('open','loop','[Từng lô tối đa 16 đoạn]'),('msg','S','O','Embedding lô với thông tin nguồn'),
  ('msg','S','Q','Upsert vector + index_version mới'),('end',),
  ('msg','S','DB','indexed + version + model; COMMIT'),
  ('else','[Trích xuất hoặc provider/vector lỗi]'),('msg','S','DB','ROLLBACK phần ghi dở; failed; COMMIT'),('end',),('end',),
  ('ret','S','C','Thông tin trạng thái/số đoạn'),('ret','C','UI','Hiển thị indexed hoặc failed'),
  ('open','opt','[Admin chọn xóa sau đó]'),('msg','UI','C','DELETE tài liệu; admin/CSRF'),
  ('msg','C','S','delete_document dưới document_lock'),
  ('msg','S','DB','deleting + version=null; COMMIT để ẩn'),
  ('msg','S','Q','Dọn vector'),('msg','S','F','Dọn file'),('msg','S','DB','Xóa đoạn/bản ghi; COMMIT'),
  ('note','Dọn dở: trả 503, nguồn vẫn bị ẩn; thử xóa lại. Reindex tạo phiên bản mới.'),('end',),
 ]),
}

class Canvas:
    def __init__(self,w,h):
        self.w,self.h=w,h
        self.img=Image.new('RGB',(w*2,h*2),'white'); self.draw=ImageDraw.Draw(self.img)
        self.svg=[f'<svg xmlns="http://www.w3.org/2000/svg" width="{w}" height="{h}" viewBox="0 0 {w} {h}"><rect width="100%" height="100%" fill="white"/>']

    def shape(self,kind,box,fill='white',stroke='#333333',width=2):
        x,y,x2,y2=box; scale=lambda vals: tuple(round(v*2) for v in vals)
        if kind=='ellipse':
            self.draw.ellipse(scale(box),fill=fill,outline=stroke,width=width*2)
            self.svg.append(f'<ellipse cx="{(x+x2)/2}" cy="{(y+y2)/2}" rx="{(x2-x)/2}" ry="{(y2-y)/2}" fill="{fill}" stroke="{stroke}" stroke-width="{width}"/>')
        else:
            self.draw.rectangle(scale(box),fill=fill,outline=stroke,width=width*2)
            self.svg.append(f'<rect x="{x}" y="{y}" width="{x2-x}" height="{y2-y}" fill="{fill}" stroke="{stroke}" stroke-width="{width}"/>')

    def line(self,points,dashed=False,width=2):
        for a,b in zip(points,points[1:]):
            if dashed:
                distance=math.dist(a,b)
                for start in range(0,math.ceil(distance),12):
                    aa=tuple(a[i]+(b[i]-a[i])*start/distance for i in (0,1))
                    bb=tuple(a[i]+(b[i]-a[i])*min(start+6,distance)/distance for i in (0,1))
                    self.draw.line(tuple(v*2 for p in (aa,bb) for v in p),fill='#333333',width=width*2)
            else: self.draw.line(tuple(v*2 for p in (a,b) for v in p),fill='#333333',width=width*2)
        pts=' '.join(f'{x},{y}' for x,y in points)
        dash=' stroke-dasharray="6 6"' if dashed else ''
        self.svg.append(f'<polyline points="{pts}" fill="none" stroke="#333333" stroke-width="{width}"{dash}/>')

    def polygon(self,points,fill='#333333'):
        self.draw.polygon([(x*2,y*2) for x,y in points],fill=fill,outline='#333333',width=4)
        self.svg.append('<polygon points="'+' '.join(f'{x},{y}' for x,y in points)+f'" fill="{fill}" stroke="#333333" stroke-width="2"/>')

    def text(self,text,x,y,size=28,max_width=None,bold=False,background=False):
        font=ImageFont.truetype('C:/Windows/Fonts/'+('arialbd.ttf' if bold else 'arial.ttf'),size*2)
        lines=[]
        for raw in text.split('\n'):
            line=''
            for word in raw.split():
                candidate=(line+' '+word).strip()
                if max_width and line and self.draw.textlength(candidate,font=font)>max_width*2:
                    lines.append(line); line=word
                else: line=candidate
            lines.append(line)
        height=(size+6)*len(lines)
        for j,line in enumerate(lines):
            length=self.draw.textlength(line,font=font)/2
            assert 0<=x-length/2 and x+length/2<=self.w, (line,x,length)
            yy=y+(size+6)*j
            if background: self.shape('rect',(x-length/2-4,yy-2,x+length/2+4,yy+size+3),'white','white',1)
            self.draw.text((x*2,yy*2),line,font=font,fill='#111111',anchor='mt')
            self.svg.append(f'<text x="{x}" y="{yy+size}" text-anchor="middle" font-family="Arial" font-size="{size}" font-weight="{"bold" if bold else "normal"}" fill="#111111">{escape(line)}</text>')
        return height

    def save(self,name,title,puml):
        self.svg.insert(1,f'<title>{escape(title)}</title>')
        (OUT/(name+'.svg')).write_text('\n'.join(self.svg+['</svg>']),encoding='utf-8')
        self.img.save(OUT/(name+'.png'))
        (OUT/(name+'.puml')).write_text('@startuml\nskinparam monochrome true\nskinparam defaultFontName Arial\ntitle '+title+'\n'+puml+'\n@enduml\n',encoding='utf-8')

def actor(c,id,x,y):
    if id in ('TG','SMTP'):
        c.shape('rect',(x-110,y-30,x+110,y+70),'#F4F5F6')
        c.text('<<actor>>\n'+ACTORS[id],x,y-22,size=25,max_width=210)
    else:
        c.shape('ellipse',(x-18,y-30,x+18,y+6))
        c.line([(x,y+6),(x,y+66)])
        c.line([(x-34,y+28),(x+34,y+28)])
        c.line([(x-30,y+100),(x,y+66),(x+30,y+100)])
        c.text(ACTORS[id],x,y+105,size=27,max_width=230,bold=True)

def usecases(name,title,ids,overall=False):
    if overall:
        c=Canvas(2400,1700); boundary=(270,95,2130,1370)
        locations={}
        groups=[(['UC03','UC04','UC06','UC10','UC13','UC15'],650,225,175),
                (['UC02','UC07','UC08','UC09','UC11','UC14','UC16','UC18'],1230,205,145),
                (['UC01','UC05','UC19','UC20','UC12','UC17'],1800,225,175)]
        for group,x,y,step in groups:
            for i,id in enumerate(group): locations[id]=(x,y+i*step)
        actors={'KH':(130,230),'TG':(130,650),'SMTP':(130,1070),'AD':(2260,225),
                'NV':(1240,1410),'RV':(2260,930),'OP':(2260,1170)}
    else:
        h=max(740,310+len(ids)*150)
        c=Canvas(1800,h); boundary=(310,100,1490,h-100)
        locations={id:(900,195+i*150) for i,id in enumerate(ids)}
        used=[a for a,goals in ASSOCIATIONS.items() if set(goals)&set(ids)]
        left=[a for a in used if a not in ('TG','SMTP','RV','OP')]
        right=[a for a in used if a in ('TG','SMTP','RV','OP')]
        actors={a:(150,180+i*min(250,(h-390)/max(1,len(left)-1))) for i,a in enumerate(left)}
        actors.update({a:(1650,180+i*min(250,(h-390)/max(1,len(right)-1))) for i,a in enumerate(right)})
    c.text(title,c.w/2,20,size=34,bold=True)
    c.shape('rect',boundary)
    c.text('Nền tảng CSKH' if overall else title,c.w/2,110,size=26,bold=True)
    # Route associations through clear rails outside ellipse interiors.
    for a,(ax,ay) in actors.items():
        for id in ASSOCIATIONS[a]:
            if id not in locations: continue
            x,y=locations[id]; rx=235 if overall else 270
            if a=='NV' and overall:
                points=[(ax,ay+5),(1530,ay+5),(1530,y+60),(x+235,y+60)]
            elif a=='TG' and overall and x==1230:
                points=[(ax+110,ay+35),(330,ay+35),(330,1355),(955,1355),(955,y+60),(x-235,y+60)]
            elif ax<c.w/2:
                rail=boundary[0]+25+list(actors).index(a)*18
                points=[(ax+40 if a not in ('TG','SMTP') else ax+110,ay+35),(rail,ay+35),(rail,y+60),(x-rx,y+60)]
            else:
                rail=boundary[2]-25-list(actors).index(a)*12
                points=[(ax-40 if a not in ('TG','SMTP') else ax-110,ay+35),(rail,ay+35),(rail,y+60),(x+rx,y+60)]
            c.line(points,width=2)
    if overall:
        # UML generalization: hollow triangle points at the general actor (staff).
        c.line([(2290,325),(2360,325),(2360,1450),(1270,1450)])
        c.polygon([(1270,1438),(1270,1462),(1240,1450)],'white')
        c.text('Admin kế thừa vai trò nhân viên; vẫn kiểm tra người phụ trách',c.w/2,1590,size=23,max_width=1800,background=True)
    for id,(x,y) in locations.items():
        rx=235 if overall else 270
        c.shape('ellipse',(x-rx,y,x+rx,y+120),'#F4F5F6')
        c.text(id+'\n'+UC[id],x,y+10,size=27,max_width=rx*2-45)
    for a,(x,y) in actors.items(): actor(c,a,x,y)
    c.text('Association: tham gia mục tiêu | Quyền cụ thể theo đặc tả | Không dùng include/extend',c.w/2,c.h-44,size=24)
    puml=[]
    for a in actors: puml.append(f'actor "{ACTORS[a]}" as {a}')
    if overall: puml.append('AD --|> NV')
    puml.append('rectangle "Nền tảng CSKH" {')
    for id in ids: puml.append(f'  usecase "{id}\\n{UC[id]}" as {id}')
    puml.append('}')
    for a in actors:
        for id in ASSOCIATIONS[a]:
            if id in ids: puml.append(f'{a} -- {id}')
    c.save(name,title,'\n'.join(puml))

def sequence(name,title,participants,events):
    w=max(1800,len(participants)*285)
    # Measure event text to reserve enough space; no fixed clipping of wrapped labels.
    heights=[]
    for e in events:
        if e[0] in ('msg','ret'):
            heights.append(88 if len(e[3])>65 else 56)
        elif e[0]=='note': heights.append(88)
        elif e[0]=='end': heights.append(24)
        else: heights.append(40)
    h=315+sum(heights)
    c=Canvas(w,h); c.text(name+' '+title,w/2,18,size=34,bold=True)
    xs={a:140+i*(w-280)/(len(participants)-1) for i,(a,_,_) in enumerate(participants)}
    for a,label,role in participants:
        x=xs[a]
        c.line([(x,195),(x,h-85)],True)
        c.shape('rect',(x-125,85,x+125,190),'#E8EDF2')
        c.text(role+'\n'+label,x,92,size=25,max_width=240,bold=True)
    y=215; stack=[]; puml=[f'participant "{label}\\n({role})" as {a}' for a,label,role in participants]; number=0
    for e,height in zip(events,heights):
        kind=e[0]
        if kind in ('msg','ret'):
            number+=1; a,b,label=e[1:]; x1,x2=xs[a],xs[b]
            direction=1 if x2>x1 else -1
            if a==b:
                c.line([(x1,y+40),(x1+70,y+40),(x1+70,y+65),(x1,y+65)],kind=='ret')
                point=(x1,y+65); direction=-1
            else:
                c.line([(x1,y+height-10),(x2,y+height-10)],kind=='ret'); point=(x2,y+height-10)
            xx,yy=point
            if kind=='ret': c.line([(xx-direction*13,yy-7),(xx,yy),(xx-direction*13,yy+7)])
            else: c.polygon([(xx,yy),(xx-direction*14,yy-7),(xx-direction*14,yy+7)])
            center=(x1+x2)/2
            center=max(460,min(w-460,center))
            c.text(f'{number}. '+label,center,y+2,size=30,max_width=900,background=True)
            puml.append(f'{a} {"-->" if kind=="ret" else "->"} {b}: {number}. {label}')
        elif kind=='note':
            c.shape('rect',(45,y+2,w-45,y+height-8),'#F4F5F6')
            c.text(e[1],w/2,y+12,size=27,max_width=w-130)
            puml.extend([f'note over {participants[0][0]}, {participants[-1][0]}',e[1],'end note'])
        elif kind=='open':
            margin=25+len(stack)*18
            stack.append((margin,y,e[1]))
            c.line([(margin,y),(w-margin,y)])
            c.line([(margin+130,y),(margin+130,y+24),(margin+115,y+36),(margin,y+36)])
            c.text(e[1],margin+65,y+3,size=26,bold=True,background=True)
            c.text(e[2],w/2,y+3,size=26,max_width=w-350,bold=True,background=True)
            puml.append(e[1]+' '+e[2])
        elif kind=='else':
            margin=stack[-1][0]; c.line([(margin,y),(w-margin,y)],True)
            c.text(e[1],w/2,y+3,size=26,max_width=w-160,bold=True,background=True); puml.append('else '+e[1])
        elif kind=='end':
            margin,top,_=stack.pop(); c.line([(margin,top),(margin,y+22),(w-margin,y+22),(w-margin,top)])
            puml.append('end')
        y+=height
    assert not stack
    c.text('Thời gian từ trên xuống | Nét đứt: phản hồi | Khung có điều kiện: nhánh/lặp/vùng ghi',w/2,h-40,size=24)
    c.save(name,title,'\n'.join(puml))

def report_builder():
    spec=importlib.util.spec_from_file_location('week04_build',ROOT/'docs/week04/build.py')
    module=importlib.util.module_from_spec(spec); spec.loader.exec_module(module)
    return module

def sync():
    path=ROOT/'docs/report/BaoCao.md'; old=path.read_text(encoding='utf-8')
    source=(HERE/'README.md').read_text(encoding='utf-8')
    ch3=source[source.index('### 3.10'):source.index('## Nội dung bổ sung Chương 4')].strip()
    ch4=source[source.index('### 4.6'):source.index('## Sử dụng hồ sơ')].strip()
    ch3=ch3.replace('](diagrams/','](../week05/diagrams/')
    ch4=ch4.replace('](diagrams/','](../week05/diagrams/')
    before=old[:old.index('## Chương 4')]
    if '### 3.10 ' in before: before=before[:before.index('### 3.10 ')]
    middle=old[old.index('## Chương 4'):old.index('## Chương 5')]
    if '### 4.6 ' in middle: middle=middle[:middle.index('### 4.6 ')]
    updated=before.rstrip()+'\n\n'+ch3+'\n\n'+middle.rstrip()+'\n\n'+ch4+'\n\n'+old[old.index('## Chương 5'):]
    updated=updated.replace('Báo cáo tiến độ và hồ sơ tuần 04 cập nhật','Báo cáo tiến độ và hồ sơ tuần 04 và 05 cập nhật')
    updated=updated.replace('Bộ hình BFD và BPMN được đặt cùng nhau ở cuối Chương 4','Bộ hình BFD và BPMN được đặt cùng nhau sau mục 4.5')
    path.write_text(updated,encoding='utf-8')
    report_builder().make_docx(updated)

def check():
    assert set(UC)=={f'UC{i:02}' for i in range(1,21)}
    assert set(UC)==set(id for _,ids in MODULES.values() for id in ids)
    assert all(id in UC for ids in ASSOCIATIONS.values() for id in ids)
    for name,(title,participants,events) in SEQUENCES.items():
        ids=[a for a,_,_ in participants]; assert len(ids)==len(set(ids))
        stack=[]
        for e in events:
            if e[0] in ('msg','ret'): assert e[1] in ids and e[2] in ids and e[3]
            elif e[0]=='open': stack.append([e[1],0])
            elif e[0]=='else': assert stack and stack[-1][0]=='alt'; stack[-1][1]+=1
            elif e[0]=='end':
                kind,branches=stack.pop()
                if kind=='alt': assert branches>=1
        assert not stack
    names=['UC_ALL',*(f'UC_{m}' for m in MODULES),*SEQUENCES]
    assert len(names)==12
    for name in names:
        ET.parse(OUT/(name+'.svg'))
        with Image.open(OUT/(name+'.png')) as im: assert im.width>=3000 and im.height>=1000
        puml=(OUT/(name+'.puml')).read_text(encoding='utf-8')
        assert puml.startswith('@startuml\n') and puml.endswith('@enduml\n')
        if name.startswith('UC_'):
            assert 'usecase ' in puml and '<<include>>' not in puml and '<<extend>>' not in puml
    report=(ROOT/'docs/report/BaoCao.md').read_text(encoding='utf-8')
    assert report.count('## Chương ')==7 and report.count('### 3.10 ')==1 and report.count('### 4.6 ')==1
    assert report.count('| Trường | Đặc tả |')==10
    for prefix,count in [('UC',20),('FR',14),('NFR',10),('BR',12)]:
        for i in range(1,count+1): assert f'{prefix}{i:02}' in report
    for image in re.findall(r'!\[[^\]]*\]\(([^)]+)\)',report): assert (ROOT/'docs/report'/image).is_file()
    from docx import Document
    d=Document(ROOT/'docs/report/BaoCao.docx')
    assert len(d.inline_shapes)==17 and len(d.tables)>=23
    text='\n'.join(p.text for p in d.paragraphs)
    assert sum(p.text.startswith('Chương ') for p in d.paragraphs)==7
    assert '3.11.10 UC10' in text and '4.6.7 Đối chiếu' in text
    print('OK: 20 Use Cases, 6 modules, 10 specifications, 7 Use Case figures, 5 Sequence figures; 7 report chapters and 17 embedded figures.')

if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__); parser.add_argument('--check',action='store_true')
    args=parser.parse_args()
    if not args.check:
        OUT.mkdir(parents=True,exist_ok=True)
        usecases('UC_ALL','Use Case toàn hệ thống',list(UC),True)
        for m,(title,ids) in MODULES.items(): usecases('UC_'+m,m+' '+title,ids)
        for name,(title,participants,events) in SEQUENCES.items(): sequence(name,title,participants,events)
        sync()
    check()
