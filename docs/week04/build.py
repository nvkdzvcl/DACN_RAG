"""Build Week 04 models and report; --check leaves one runnable model check.

ponytail: fixed layouts for four analysis models; use BPMN editor for larger flows.
"""
import argparse
import re
from pathlib import Path
from xml.etree import ElementTree as ET

ROOT = Path(__file__).resolve().parents[2]
HERE = Path(__file__).resolve().parent
OUT = HERE / 'diagrams'
NS = {
    'bpmn': 'http://www.omg.org/spec/BPMN/20100524/MODEL',
    'bpmndi': 'http://www.omg.org/spec/BPMN/20100524/DI',
    'dc': 'http://www.omg.org/spec/DD/20100524/DC',
    'di': 'http://www.omg.org/spec/DD/20100524/DI',
    'xsi': 'http://www.w3.org/2001/XMLSchema-instance',
}
for prefix, uri in NS.items():
    ET.register_namespace(prefix, uri)


def xml(parent, tag, **attrs):
    prefix, local = tag.split(':')
    return ET.SubElement(parent, '{%s}%s' % (NS[prefix], local), {k: str(v) for k, v in attrs.items()})


class Model:
    def __init__(self, name, title, lanes, height, external=()):
        self.name, self.title, self.height = name, title, height
        self.lanes = lanes  # (id, title, y, height)
        self.external = external  # (id, title, y, height)
        self.nodes, self.flows, self.messages = {}, [], []

    def node(self, id, kind, label, lane, x, y, w=180, h=76, **attrs):
        if kind.endswith('Event'):
            w = h = 42 if kind != 'boundaryEvent' else 30
        elif kind == 'exclusiveGateway':
            w = h = 58
        self.nodes[id] = dict(id=id, kind=kind, label=label, lane=lane, x=x, y=y, w=w, h=h, **attrs)
        return id

    def center(self, id):
        n = self.nodes[id]
        return n['x'] + n['w']/2, n['y'] + n['h']/2

    def flow(self, source, target, label='', via=None, default=False, label_at=None):
        a, b = self.nodes[source], self.nodes[target]
        ax, ay = self.center(source)
        bx, by = self.center(target)
        if via:
            # Explicit first/last point controls the attached side.
            points = via
        elif abs(ay-by) < 100 and (a['x']+a['w']<b['x'] or b['x']+b['w']<a['x']):
            sx=a['x']+a['w'] if bx>ax else a['x']
            tx=b['x'] if bx>ax else b['x']+b['w']
            middle=(sx+tx)/2
            points=[(sx,ay),(middle,ay),(middle,by),(tx,by)]
        elif abs(ax-bx) < 10:
            points = [(ax, a['y']+a['h']), (bx, b['y'])] if by > ay else [(ax, a['y']), (bx, b['y']+b['h'])]
        else:
            sy=a['y']+a['h'] if by>ay else a['y']
            ty=b['y'] if by>ay else b['y']+b['h']
            middle=(sy+ty)/2
            points=[(ax,sy),(ax,middle),(bx,middle),(bx,ty)]
        self.flows.append(dict(id=f'F_{len(self.flows)+1}', source=source, target=target,
                               label=label, points=points, default=default, label_at=label_at))

    def message(self, pool, node, label, points, outbound=False, label_at=None):
        self.messages.append(dict(id=f'M_{len(self.messages)+1}', source=node if outbound else pool,
                                 target=pool if outbound else node, label=label, points=points,label_at=label_at))


def models():
    m = Model('BPMN01', 'Tiếp nhận hỏi chính sách và chuyển giao',
              [('app', 'Ứng dụng FastAPI', 165, 535), ('rag', 'Ollama và Qdrant local', 700, 180)], 930,
              [('customer', 'Khách hàng', 45, 65)])
    n = m.node
    n('start', 'startEvent', 'Nhận yêu cầu', 'app', 160, 208, message=True)
    n('validate', 'serviceTask', 'Kiểm tra phiên và\ndữ liệu đầu vào', 'app', 245, 190)
    n('valid', 'exclusiveGateway', 'Hợp lệ?', 'app', 475, 200)
    n('save', 'serviceTask', 'Lưu tin khách;\nkiểm tra UUID', 'app', 580, 190)
    n('duplicate', 'exclusiveGateway', 'Tin trùng?', 'app', 815, 200)
    n('route', 'exclusiveGateway', 'Trạng thái và\nyêu cầu?', 'app', 1060, 200)
    n('reject', 'serviceTask', 'Báo lỗi; giữ\nbản nháp tại UI', 'app', 320, 355)
    n('end_invalid', 'endEvent', 'Yêu cầu bị chặn', 'app', 590, 372)
    n('reuse', 'serviceTask', 'Trả kết quả đã lưu;\nkhông gọi lại AI', 'app', 710, 355)
    n('end_duplicate', 'endEvent', 'Đã phục hồi', 'app', 955, 372)
    n('order', 'serviceTask', 'Tra đơn theo\nBPMN02', 'app', 555, 450)
    n('end_order', 'endEvent', 'Kết thúc lượt tra', 'app', 800, 467)
    n('queue', 'serviceTask', 'Tạo hoặc tái dùng\nticket; dừng AI', 'app', 1210, 350)
    n('end_queue', 'endEvent', 'Chờ BPMN03', 'app', 1490, 367)
    n('keep', 'serviceTask', 'Chỉ lưu tin; chờ\nnhân viên xử lý', 'app', 1210, 495)
    n('end_wait', 'endEvent', 'AI không chạy', 'app', 1490, 512)
    n('call_rag', 'serviceTask', 'Truy hồi nguồn;\nchọn và kiểm định', 'rag', 810, 750)
    n('rag_error', 'boundaryEvent', 'Lỗi dịch vụ', 'rag', 890, 811, attachedToRef='call_rag', error=True)
    n('error_notice', 'serviceTask', 'Báo lỗi dịch vụ;\ngiữ tin khách', 'app', 555, 575)
    n('end_error', 'endEvent', 'Chưa có đáp án', 'app', 330, 592)
    n('guard', 'exclusiveGateway', 'Trạng thái, tin cuối,\nnguồn còn đúng?', 'app', 1010, 575)
    n('drop', 'serviceTask', 'Bỏ kết quả AI cũ', 'app', 1210, 595)
    n('end_stale', 'endEvent', 'Không gửi AI cũ', 'app', 1490, 612)
    n('answer', 'sendTask', 'Trả bằng chứng hoặc\ntừ chối / hỏi rõ', 'app', 1300, 190)
    n('end_ok', 'endEvent', 'Kết thúc lượt', 'app', 1560, 208)
    m.flow('start','validate'); m.flow('validate','valid')
    m.flow('valid','save','Có')
    m.flow('valid','reject','Không',[(504,258),(504,325),(410,325),(410,355)],default=True)
    m.flow('reject','end_invalid'); m.flow('save','duplicate')
    m.flow('duplicate','reuse','Có',[(844,258),(844,325),(800,325),(800,355)],label_at=(800,338))
    m.flow('reuse','end_duplicate'); m.flow('duplicate','route','Không',default=True)
    m.flow('route','call_rag','Open; chính sách',[(1089,258),(1089,315),(1120,315),(1120,720),(900,720),(900,750)],label_at=(1030,704))
    m.flow('route','order','Open; tra đơn',[(1089,258),(1089,438),(645,438),(645,450)],label_at=(1020,350))
    m.flow('order','end_order')
    m.flow('route','queue','Cần người thật',[(1118,229),(1160,229),(1160,388),(1210,388)],label_at=(1280,318))
    m.flow('route','keep','Chờ / assigned',[(1089,258),(1145,258),(1145,533),(1210,533)],label_at=(1260,467))
    m.flow('route','reject','Closed',[(1089,200),(1089,175),(130,175),(130,393),(320,393)],default=True,label_at=(900,161))
    m.flow('queue','end_queue'); m.flow('keep','end_wait')
    m.flow('call_rag','guard','',[(990,788),(1039,788),(1039,633)])
    m.flow('rag_error','error_notice','',[(890,826),(780,826),(780,613),(735,613)])
    m.flow('error_notice','end_error')
    m.flow('guard','drop','Không',default=True); m.flow('drop','end_stale')
    m.flow('guard','answer','Có',[(1010,604),(1000,604),(1000,305),(1390,305),(1390,266)],label_at=(1370,285))
    m.flow('answer','end_ok')
    m.message('customer','start','Tin khách',[(181,110),(181,208)])
    m.message('customer','answer','Phản hồi qua kênh',[(1390,190),(1390,110)],outbound=True)

    o = Model('BPMN02', 'Xác minh và tra cứu đơn hàng',
              [('app','Ứng dụng FastAPI',165,420), ('admin','Quản trị viên',585,270)], 910,
              [('customer','Khách hàng widget',45,65)])
    n=o.node
    n('start','startEvent','Yêu cầu\ntra đơn','app',155,202,message=True)
    n('scope','exclusiveGateway','Widget và chỉ\nyêu cầu tra cứu?','app',260,194)
    n('access','exclusiveGateway','Đã có quyền\ncòn hiệu lực?','app',465,194)
    n('verify','userTask','Xác minh qua kênh\nliên hệ tin cậy','admin',350,650)
    n('verified','exclusiveGateway','Xác minh đạt?','admin',625,659)
    n('issue','serviceTask','Cấp mã một đơn;\nhạn 15 phút','app',755,385)
    n('deliver','userTask','Giao mã qua kênh\nđã xác minh','admin',755,650)
    n('receive','receiveTask','Nhận mã tại\nbiểu mẫu riêng','app',1040,385)
    n('code','exclusiveGateway','Mã và quyền\nhợp lệ?','app',1300,394,label_at=(1319,488))
    n('lookup','serviceTask','Tra trạng thái và\nvận đơn được phép','app',755,194)
    n('guard','exclusiveGateway','Quyền và chủ đơn\ncòn đúng?','app',1040,203)
    n('reply','sendTask','Chỉ trả trạng thái\nvà vận đơn','app',1200,194)
    n('end_ok','endEvent','Kết thúc lượt','app',1500,211)
    n('handoff','serviceTask','Handoff; không lộ\ndữ liệu đơn','app',1380,470,w=160)
    n('end_queue','endEvent','Chờ xử lý','app',1590,487)
    o.flow('start','scope')
    o.flow('scope','access','Có')
    o.flow('scope','handoff','Telegram / hủy / đổi trả',[(289,194),(289,177),(1600,177),(1600,450),(1460,450),(1460,470)],default=True,label_at=(830,163))
    o.flow('access','lookup','Có')
    o.flow('access','verify','Chưa',[(494,252),(494,305),(440,305),(440,650)],default=True)
    o.flow('verify','verified')
    o.flow('verified','issue','Có',[(683,688),(715,688),(715,423),(755,423)],label_at=(738,573))
    o.flow('verified','handoff','Không',[(654,717),(654,795),(1460,795),(1460,546)],default=True)
    o.flow('issue','deliver'); o.flow('deliver','receive','',[(935,688),(985,688),(985,423),(1040,423)])
    o.flow('receive','code')
    o.flow('code','lookup','Có',[(1329,394),(1329,325),(845,325),(845,270)],label_at=(870,310))
    o.flow('code','handoff','Sai / hết hạn',[(1358,423),(1370,423),(1370,508),(1380,508)],default=True,label_at=(1460,413))
    o.flow('lookup','guard'); o.flow('guard','reply','Có'); o.flow('reply','end_ok')
    o.flow('guard','handoff','Thu hồi / hết hạn',[(1069,261),(1069,295),(1550,295),(1550,460),(1485,460),(1485,470)],default=True,label_at=(1390,283))
    o.flow('handoff','end_queue')
    o.message('customer','start','Yêu cầu',[(176,110),(176,202)])
    o.message('customer','deliver','Mã qua kênh tin cậy',[(845,650),(700,650),(700,120),(700,110)],outbound=True)
    o.message('customer','receive','Nhập mã tại biểu mẫu',[(1130,110),(1130,385)],label_at=(1120,143))
    o.message('customer','reply','Kết quả được phép',[(1290,194),(1290,110)],outbound=True)

    t = Model('BPMN03','Nhân viên xử lý ticket và giao phản hồi',
              [('staff','Nhân viên phụ trách',145,290),('app','Ứng dụng FastAPI và worker',435,520)], 1100,
              [('customer','Khách hàng',40,60),('telegram','Telegram Bot API',995,60)])
    n=t.node
    n('start','startEvent','Có ticket chờ','staff',145,192)
    n('accept','userTask','Đọc hàng chờ;\nyêu cầu tiếp nhận','staff',235,175)
    n('assign','serviceTask','Gán có điều kiện;\nkiểm tra phụ trách','app',235,480)
    n('assigned','exclusiveGateway','Nhận thành công?','app',460,489)
    n('conflict','serviceTask','Báo tranh chấp;\nyêu cầu tải lại','app',140,720)
    n('end_conflict','endEvent','Người khác nhận','app',200,875)
    n('compose','userTask','Đọc tin; soạn\nphản hồi / đọc lại','staff',575,175)
    n('save','serviceTask','Lưu phản hồi staff\nvới UUID','app',575,480)
    n('channel','exclusiveGateway','Kênh giao tin?','app',800,489)
    n('web','sendTask','Website: cho khách\nđọc bằng polling','app',920,480)
    n('tg','serviceTask','Telegram: hàng gửi\nvà sendMessage','app',760,720)
    n('tg_error','boundaryEvent','Lỗi giao tin','app',840,781,attachedToRef='tg',error=True,label_at=(772,845))
    n('failed','serviceTask','Ghi lỗi / chưa rõ;\nchặn tin cùng peer','app',1040,785)
    n('decision','userTask','Retry hoặc skip;\nuncertain: xác nhận','staff',1040,320)
    n('retry','exclusiveGateway','Gửi lại?','app',1280,793)
    n('skip','serviceTask','Giữ lịch sử;\nđánh dấu skipped','app',1400,785)
    n('more','exclusiveGateway','Hỗ trợ tiếp?','staff',1235,184)
    n('finish','userTask','Resolved / closed;\nghi chú và tin cuối','staff',1370,175)
    n('check','serviceTask','Kiểm tra ticket,\nngười và tin cuối','app',1370,480)
    n('current','exclusiveGateway','Đủ điều kiện?','app',1400,630)
    n('commit','serviceTask','Chốt ticket và SLA;\nlưu người, thời điểm','app',1130,635)
    n('end_ok','endEvent','Kết thúc lượt hỗ trợ','app',1040,651)
    t.flow('start','accept'); t.flow('accept','assign'); t.flow('assign','assigned')
    t.flow('assigned','compose','Có',[(518,518),(540,518),(540,213),(575,213)])
    t.flow('assigned','conflict','Không',[(489,547),(489,690),(230,690),(230,720)],default=True)
    t.flow('conflict','end_conflict','',[(230,796),(230,835),(221,835),(221,875)])
    t.flow('compose','save'); t.flow('save','channel')
    t.flow('channel','web','Website'); t.flow('web','more','',[(1100,518),(1250,518),(1250,280),(1264,280),(1264,242)])
    t.flow('channel','tg','Telegram',[(829,547),(829,675),(850,675),(850,720)],default=True)
    t.flow('tg','more','Đã gửi',[(940,758),(1320,758),(1320,155),(1264,155),(1264,184)],label_at=(1030,744))
    t.flow('tg_error','failed','',[(870,796),(940,796),(940,823),(1040,823)])
    t.flow('failed','decision','',[(1130,785),(1130,396)])
    t.flow('decision','retry','',[(1220,358),(1310,358),(1310,770),(1309,770),(1309,793)])
    t.flow('retry','tg','Có, xác nhận khi cần',[(1309,851),(1309,920),(670,920),(670,758),(760,758)])
    t.flow('retry','skip','Không',default=True)
    t.flow('skip','more','',[(1580,823),(1620,823),(1620,420),(1264,420),(1264,242)])
    t.flow('more','compose','Có',[(1235,213),(1190,213),(1190,275),(665,275),(665,251)],label_at=(845,260))
    t.flow('more','finish','Hoàn tất',default=True,label_at=(1330,165))
    t.flow('finish','check'); t.flow('check','current')
    t.flow('current','commit','Có'); t.flow('commit','end_ok')
    t.flow('current','compose','Xung đột; đọc lại',[(1429,688),(1429,740),(1340,740),(1340,300),(665,300),(665,251)],default=True,label_at=(850,319))
    t.message('telegram','tg','sendMessage',[(795,796),(795,995)],outbound=True,label_at=(795,946))
    t.message('telegram','tg','Kết quả gửi',[(975,995),(975,900),(915,900),(915,796)],label_at=(962,947))
    t.message('customer','web','Phản hồi Website',[(1010,480),(1010,110),(1010,100)],outbound=True,label_at=(1100,122))

    k = Model('BPMN04','Quản lý và lập chỉ mục tri thức',
              [('admin','Quản trị viên',55,225),('app','Ứng dụng FastAPI',280,375),('rag','Ollama và Qdrant local',655,235)], 940)
    n=k.node
    n('start','startEvent','Có tài liệu','admin',140,120)
    n('upload','userTask','Chọn tài liệu và\nyêu cầu tải lên','admin',235,103)
    n('validate','serviceTask','Kiểm tra loại,\n10 MB và hash','app',235,335)
    n('valid','exclusiveGateway','File hợp lệ?','app',470,344)
    n('reject','serviceTask','Từ chối; thông báo\nlỗi hoặc trùng file','app',235,525)
    n('end_invalid','endEvent','Chưa lập chỉ mục','app',470,542)
    n('extract','serviceTask','Lưu nguồn; đọc và\nchia đoạn có vị trí','app',625,335)
    n('parse_error','boundaryEvent','Lỗi đọc file','app',710,396,attachedToRef='extract',error=True)
    n('content','exclusiveGateway','Có nội dung đọc\nđược?', 'app',880,344)
    n('index','serviceTask','Tạo embedding;\nghi vector mới','rag',850,720)
    n('index_error','boundaryEvent','Lỗi chỉ mục','rag',930,781,attachedToRef='index',error=True)
    n('error','serviceTask','Đánh dấu lỗi;\ncho phép lập lại','app',625,525)
    n('end_error','endEvent','Chưa sẵn sàng','app',860,542)
    n('ready','serviceTask','Đánh dấu sẵn sàng;\ncho phép truy hồi','app',1110,335)
    n('action','exclusiveGateway','Thao tác tiếp?','admin',1300,112)
    n('reindex','userTask','Lập chỉ mục lại\nphiên bản mới','admin',850,103)
    n('delete','serviceTask','Ngừng truy hồi rồi\ndọn file / vector','app',1320,525)
    n('end_deleted','endEvent','Đã ngừng dùng nguồn','app',1550,542)
    n('end_ready','endEvent','Giữ nguồn sẵn sàng','admin',1510,120)
    k.flow('start','upload'); k.flow('upload','validate'); k.flow('validate','valid')
    k.flow('valid','extract','Có')
    k.flow('valid','reject','Không',[(499,402),(499,480),(325,480),(325,525)],default=True)
    k.flow('reject','end_invalid'); k.flow('extract','content')
    k.flow('parse_error','error')
    k.flow('content','index','Có',[(909,402),(1040,402),(1040,680),(940,680),(940,720)])
    k.flow('content','error','Không; scan chưa OCR',[(909,402),(909,480),(780,480),(780,525)],default=True,label_at=(930,490))
    k.flow('error','end_error')
    k.flow('index_error','error','',[(930,796),(575,796),(575,563),(625,563)])
    k.flow('index','ready','',[(1030,758),(1200,758),(1200,411)])
    k.flow('ready','action','',[(1290,373),(1329,373),(1329,170)])
    k.flow('action','end_ready','Giữ',default=True)
    k.flow('action','reindex','Lập lại',[(1300,141),(1030,141)])
    k.flow('reindex','extract','',[(940,179),(940,235),(715,235),(715,335)])
    k.flow('action','delete','Xóa',[(1329,170),(1585,170),(1585,480),(1410,480),(1410,525)])
    k.flow('delete','end_deleted')
    return [m,o,t,k]


def write_bpmn(m):
    root=ET.Element('{%s}definitions'%NS['bpmn'], id=m.name+'_Definitions',
                    targetNamespace='https://dacn-rag.local/week04', exporter='DACN RAG Week04', exporterVersion='1')
    xml(root,'bpmn:error',id='Error_Service',name='Lỗi dịch vụ',errorCode='service_error')
    process=xml(root,'bpmn:process',id='Process',name=m.title,isExecutable='false')
    lanes=xml(process,'bpmn:laneSet',id='LaneSet')
    for id,title,y,h in m.lanes:
        lane=xml(lanes,'bpmn:lane',id='Lane_'+id,name=title)
        for node in m.nodes.values():
            if node['lane']==id:
                xml(lane,'bpmn:flowNodeRef').text=node['id']
    for n in m.nodes.values():
        attrs=dict(id=n['id'],name=n['label'].replace('\n',' '))
        if n['kind']=='boundaryEvent':
            attrs.update(attachedToRef=n['attachedToRef'],cancelActivity='true')
        defaults=[f['id'] for f in m.flows if f['source']==n['id'] and f['default']]
        if defaults: attrs['default']=defaults[0]
        node=xml(process,'bpmn:'+n['kind'],**attrs)
        for f in m.flows:
            if f['target']==n['id']: xml(node,'bpmn:incoming').text=f['id']
        for f in m.flows:
            if f['source']==n['id']: xml(node,'bpmn:outgoing').text=f['id']
        if n.get('message'): xml(node,'bpmn:messageEventDefinition',id=n['id']+'_MessageEvent')
        if n.get('error'): xml(node,'bpmn:errorEventDefinition',id=n['id']+'_ErrorEvent',errorRef='Error_Service')
    for f in m.flows:
        flow=xml(process,'bpmn:sequenceFlow',id=f['id'],sourceRef=f['source'],targetRef=f['target'],name=f['label'])
        if m.nodes[f['source']]['kind']=='exclusiveGateway' and f['label'] and not f['default']:
            expression=xml(flow,'bpmn:conditionExpression',**{'{%s}type'%NS['xsi']:'bpmn:tFormalExpression'},language='text/plain')
            expression.text=f['label']
    collab=xml(root,'bpmn:collaboration',id='Collaboration')
    xml(collab,'bpmn:participant',id='Store',name='Cửa hàng giả lập',processRef='Process')
    for id,title,y,h in m.external: xml(collab,'bpmn:participant',id=id,name=title)
    for f in m.messages:
        xml(collab,'bpmn:messageFlow',id=f['id'],sourceRef=f['source'],targetRef=f['target'],name=f['label'])
    diagram=xml(root,'bpmndi:BPMNDiagram',id=m.name+'_Diagram')
    plane=xml(diagram,'bpmndi:BPMNPlane',id='Plane',bpmnElement='Collaboration')
    def shape(id,x,y,w,h,horizontal=False,label=None):
        attrs=dict(id='Shape_'+id,bpmnElement=id)
        if horizontal: attrs['isHorizontal']='true'
        s=xml(plane,'bpmndi:BPMNShape',**attrs); xml(s,'dc:Bounds',x=x,y=y,width=w,height=h)
        if label:
            l=xml(s,'bpmndi:BPMNLabel'); xml(l,'dc:Bounds',x=label[0],y=label[1],width=label[2],height=label[3])
    top=m.lanes[0][2]; bottom=m.lanes[-1][2]+m.lanes[-1][3]
    shape('Store',45,top,1600,bottom-top,True)
    for id,title,y,h in m.lanes: shape('Lane_'+id,75,y,1570,h,True)
    for id,title,y,h in m.external: shape(id,45,y,1600,h,True)
    for n in m.nodes.values():
        label=None
        if n['kind'].endswith('Event') or n['kind']=='exclusiveGateway':
            lx,ly=n.get('label_at',(n['x']+n['w']/2,n['y']+n['h']+29))
            label=(lx-115,ly-25,230,50)
        shape(n['id'],n['x'],n['y'],n['w'],n['h'],label=label)
    for f in m.flows+m.messages:
        e=xml(plane,'bpmndi:BPMNEdge',id='Edge_'+f['id'],bpmnElement=f['id'])
        for x,y in f['points']: xml(e,'di:waypoint',x=x,y=y)
        if f['label']:
            x,y=flow_label(f); l=xml(e,'bpmndi:BPMNLabel'); xml(l,'dc:Bounds',x=x-95,y=y-12,width=190,height=26)
    ET.indent(root)
    ET.ElementTree(root).write(OUT/(m.name+'.bpmn'),encoding='utf-8',xml_declaration=True)


def flow_label(f):
    if f.get('label_at'): return f['label_at']
    points=f['points']
    a,b=max(zip(points,points[1:]),key=lambda p: abs(p[1][0]-p[0][0])+abs(p[1][1]-p[0][1]))
    return ((a[0]+b[0])/2,(a[1]+b[1])/2-14)


class Canvas:
    """Same geometry/text for SVG and high resolution PNG, no rendering service."""
    def __init__(self,width,height):
        from PIL import Image,ImageDraw,ImageFont
        self.width,self.height=width,height
        self.image=Image.new('RGB',(width*2,height*2),'white'); self.draw=ImageDraw.Draw(self.image)
        self.font_path='C:/Windows/Fonts/arial.ttf'
        self.ImageFont=ImageFont
        self.svg=ET.Element('svg',xmlns='http://www.w3.org/2000/svg',width=str(width),height=str(height),viewBox=f'0 0 {width} {height}')
        ET.SubElement(self.svg,'rect',width='100%',height='100%',fill='white')

    def rect(self,box,fill='white',radius=0,width=2):
        x,y,w,h=box
        ET.SubElement(self.svg,'rect',x=str(x),y=str(y),width=str(w),height=str(h),rx=str(radius),fill=fill,stroke='#222222',**{'stroke-width':str(width)})
        self.draw.rounded_rectangle((x*2,y*2,(x+w)*2,(y+h)*2),radius*2,fill=fill,outline='#222222',width=width*2)

    def line(self,points,dashed=False,width=2,arrow=False,open_arrow=False):
        attrs={'points':' '.join(f'{x},{y}' for x,y in points),'fill':'none','stroke':'#333333','stroke-width':str(width)}
        if dashed: attrs['stroke-dasharray']='8 6'
        ET.SubElement(self.svg,'polyline',**attrs)
        for a,b in zip(points,points[1:]):
            if dashed:
                import math
                length=math.dist(a,b)
                for distance in range(0,int(length),14):
                    end=min(distance+8,length)
                    p=(a[0]+(b[0]-a[0])*distance/length,a[1]+(b[1]-a[1])*distance/length)
                    q=(a[0]+(b[0]-a[0])*end/length,a[1]+(b[1]-a[1])*end/length)
                    self.draw.line([(p[0]*2,p[1]*2),(q[0]*2,q[1]*2)],fill='#333333',width=width*2)
            else: self.draw.line([(a[0]*2,a[1]*2),(b[0]*2,b[1]*2)],fill='#333333',width=width*2)
        if arrow:
            import math
            a,b=points[-2:]; angle=math.atan2(b[1]-a[1],b[0]-a[0]); length=12
            left=(b[0]-length*math.cos(angle-.45),b[1]-length*math.sin(angle-.45))
            right=(b[0]-length*math.cos(angle+.45),b[1]-length*math.sin(angle+.45))
            if open_arrow: self.line([left,b,right],width=width)
            else: self.polygon([left,b,right],fill='#333333')

    def polygon(self,points,fill='white',width=2):
        ET.SubElement(self.svg,'polygon',points=' '.join(f'{x},{y}' for x,y in points),fill=fill,stroke='#222222',**{'stroke-width':str(width)})
        pts=[(x*2,y*2) for x,y in points]; self.draw.polygon(pts,fill=fill); self.draw.line(pts+[pts[0]],fill='#222222',width=width*2)

    def circle(self,x,y,r,width=2,fill='white'):
        ET.SubElement(self.svg,'circle',cx=str(x),cy=str(y),r=str(r),fill=fill,stroke='#222222',**{'stroke-width':str(width)})
        self.draw.ellipse(((x-r)*2,(y-r)*2,(x+r)*2,(y+r)*2),fill=fill,outline='#222222',width=width*2)

    def text(self,text,x,y,size=20,max_width=215,fill='#111111',background=False):
        font=self.ImageFont.truetype(self.font_path,size*2)
        lines=[]
        for paragraph in text.split('\n'):
            current=''
            for word in paragraph.split():
                trial=(current+' '+word).strip()
                if current and self.draw.textlength(trial,font=font)>max_width*2:
                    lines.append(current); current=word
                else: current=trial
            lines.append(current)
        line_height=size*1.16; top=y-(len(lines)-1)*line_height/2
        if background:
            width=max(self.draw.textlength(s,font=font)/2 for s in lines)+8
            # Labels stay white without decorative outline.
            ET.SubElement(self.svg,'rect',x=str(x-width/2),y=str(top-size*.55),width=str(width),height=str(line_height*len(lines)),fill='white')
            self.draw.rectangle(((x-width/2)*2,(top-size*.55)*2,(x+width/2)*2,(top-size*.55+line_height*len(lines))*2),fill='white')
        for i,line in enumerate(lines):
            baseline=top+i*line_height
            t=ET.SubElement(self.svg,'text',x=str(x),y=str(baseline),fill=fill,**{'font-family':'Arial','font-size':str(size),'text-anchor':'middle','dominant-baseline':'central'})
            t.text=line
            self.draw.text((x*2,baseline*2),line,font=font,fill=fill,anchor='mm')

    def vertical(self,text,x,y,size=19):
        t=ET.SubElement(self.svg,'text',x=str(x),y=str(y),transform=f'rotate(-90 {x} {y})',**{'font-family':'Arial','font-size':str(size),'text-anchor':'middle'})
        t.text=text
        from PIL import Image,ImageDraw
        font=self.ImageFont.truetype(self.font_path,size*2)
        w=int(self.draw.textlength(text,font=font))+10
        layer=Image.new('RGBA',(w,size*3),(255,255,255,0)); d=ImageDraw.Draw(layer)
        d.text((w/2,size*1.5),text,font=font,fill='#111111',anchor='mm')
        layer=layer.rotate(90,expand=True)
        self.image.paste(layer,(int(x*2-layer.width/2),int(y*2-layer.height/2)),layer)

    def save(self,name):
        ET.indent(self.svg); ET.ElementTree(self.svg).write(OUT/(name+'.svg'),encoding='utf-8',xml_declaration=True)
        self.image.save(OUT/(name+'.png'),dpi=(240,240))


def draw_model(m):
    c=Canvas(1700,m.height)
    c.text(m.name+' '+m.title,850,22,24,1600)
    top=m.lanes[0][2]; bottom=m.lanes[-1][2]+m.lanes[-1][3]
    c.rect((45,top,1600,bottom-top)); c.vertical('Cửa hàng giả lập',64,(top+bottom)/2)
    for id,title,y,h in m.lanes:
        c.rect((75,y,1570,h),fill='white'); c.line([(110,y),(110,y+h)])
        c.vertical(title,96,y+h/2,18)
    for id,title,y,h in m.external:
        c.rect((45,y,1600,h),fill='#f3f4f5'); c.text(title,850,y+h/2,21,1400)
    for f in m.flows:
        c.line(f['points'],arrow=True)
        if f['default']:
            a,b=f['points'][:2]; dx,dy=b[0]-a[0],b[1]-a[1]
            import math
            length=math.hypot(dx,dy)
            if length:
                px,py=a[0]+dx/length*17,a[1]+dy/length*17
                c.line([(px-dy/length*6-dx/length*4,py+dx/length*6-dy/length*4),
                        (px+dy/length*6+dx/length*4,py-dx/length*6+dy/length*4)])
    for f in m.messages:
        c.line(f['points'],dashed=True,arrow=True,open_arrow=True)
        c.circle(*f['points'][0],3,width=1)
    for n in m.nodes.values():
        x,y,w,h=n['x'],n['y'],n['w'],n['h']; cx,cy=x+w/2,y+h/2
        if n['kind'].endswith('Event'):
            c.circle(cx,cy,w/2,4 if n['kind']=='endEvent' else 2)
            if n['kind']=='boundaryEvent':
                c.circle(cx,cy,w/2-4,1)
                c.polygon([(cx-5,cy-7),(cx,cy-2),(cx-2,cy+2),(cx+5,cy+7),(cx+1,cy),(cx+4,cy-3)],fill='#222222',width=1)
            if n.get('message'):
                c.rect((cx-11,cy-7,22,14),width=1); c.line([(cx-11,cy-7),(cx,cy+1),(cx+11,cy-7)],width=1)
            lx,ly=n.get('label_at',(cx,y+h+25))
            c.text(n['label'],lx,ly,18,230,background=True)
        elif n['kind']=='exclusiveGateway':
            c.polygon([(cx,y),(x+w,cy),(cx,y+h),(x,cy)])
            c.line([(cx-8,cy-8),(cx+8,cy+8)]); c.line([(cx+8,cy-8),(cx-8,cy+8)])
            lx,ly=n.get('label_at',(cx,y+h+29))
            c.text(n['label'],lx,ly,18,235,background=True)
        else:
            c.rect((x,y,w,h),radius=10,fill='#fafafa')
            if n['kind']=='userTask':
                c.circle(x+15,y+12,4,width=1); c.line([(x+8,y+26),(x+10,y+19),(x+20,y+19),(x+22,y+26)],width=1)
            elif n['kind'] in ('sendTask','receiveTask'):
                c.rect((x+8,y+8,20,13),fill='#333333' if n['kind']=='sendTask' else 'white',width=1)
                c.line([(x+8,y+8),(x+18,y+15),(x+28,y+8)],width=1)
            else:
                c.circle(x+17,y+16,7,width=2); c.circle(x+17,y+16,2,width=1)
                for dx,dy in ((10,0),(-10,0),(0,10),(0,-10)):
                    c.line([(x+17+dx*.65,y+16+dy*.65),(x+17+dx,y+16+dy)],width=2)
            c.text(n['label'],cx,cy+8,18,w-14)
    for f in m.flows+m.messages:
        if f['label']:
            x,y=flow_label(f); c.text(f['label'],x,y,16,205,background=True)
    c.text('Nét liền: sequence flow | Nét đứt: message flow | X: chọn một nhánh | Vạch chéo: nhánh mặc định',850,m.height-24,17,1600)
    c.save(m.name)


def draw_bfd():
    c=Canvas(1700,990)
    c.text('BFD Phân rã chức năng nền tảng hỗ trợ khách hàng',850,28,25,1600)
    c.rect((575,65,550,75),fill='#e9edf1',radius=8)
    c.text('Hỗ trợ khách hàng đa kênh\nWebsite và Telegram',850,103,24,530)
    groups=[
        ('BFD01','Truy cập và danh tính',[
            'Tài khoản và phiên\n[Khách, staff, FastAPI]',
            'Xác minh email và reset\n[Khách, FastAPI, SMTP]',
            'Cấp quyền từng đơn\n[Admin, FastAPI]']),
        ('BFD02','Kho tri thức',[
            'Upload và lập chỉ mục\n[Admin, FastAPI, RAG]',
            'Xem nguồn và vị trí\n[Staff, FastAPI]',
            'Lập lại và xóa nguồn\n[Admin, FastAPI, RAG]']),
        ('BFD03','Hội thoại đa kênh',[
            'Nhận, lưu và giao tin\n[Khách, FastAPI, Telegram]',
            'Hỏi chính sách và tra đơn\n[Khách, FastAPI, RAG]',
            'Khôi phục và lịch sử\n[Khách, staff, FastAPI]']),
        ('BFD04','Ticket và chuyển giao',[
            'Handoff và nhận việc\n[Khách, staff, FastAPI]',
            'Trả lời và hoàn tất\n[Staff phụ trách, FastAPI]',
            'SLA và lượt hỗ trợ mới\n[Staff, FastAPI]']),
        ('BFD05','Quản lý và báo cáo',[
            'Hồ sơ và đơn giả lập\n[Admin, staff, FastAPI]',
            'Tìm kiếm và phân trang\n[Staff, FastAPI]',
            'Thống kê và phân tích\n[Staff, FastAPI]']),
        ('BFD06','Đánh giá và vận hành',[
            'Phiếu và chất lượng RAG\n[Người duyệt]',
            'Đo tải và sẵn sàng\n[Người vận hành]',
            'Backup và restore\n[Người vận hành]']),
    ]
    centers=[145+i*282 for i in range(6)]
    c.line([(850,140),(850,175)]); c.line([(centers[0],175),(centers[-1],175)])
    for cx,(id,title,children) in zip(centers,groups):
        c.line([(cx,175),(cx,215)],arrow=True)
        c.rect((cx-125,215,250,100),fill='#e9edf1',radius=8); c.text(id+'\n'+title,cx,265,22,228)
        trunk=cx-120; c.line([(cx,315),(cx,340),(trunk,340),(trunk,860)])
        for i,text in enumerate(children):
            y=375+i*190; c.line([(trunk,y+66),(cx-106,y+66)],arrow=True)
            c.rect((cx-106,y,225,132),radius=8); c.text(text,cx+6,y+66,20,207)
    c.text('RAG = Ollama và Qdrant local. Nhãn trong ngoặc vuông chỉ vai trò; cây không thể hiện thứ tự thời gian.',850,958,20,1600)
    c.save('BFD')


def check(schema_dir=None):
    from PIL import Image
    for m in models():
        root=ET.parse(OUT/(m.name+'.bpmn')).getroot()
        elements={e.get('id'):e for e in root.iter() if e.get('id')}
        ids=[e.get('id') for e in root.iter() if e.get('id')]
        assert len(ids)==len(set(ids)), f'{m.name}: duplicate ID'
        assert root.find('bpmn:process',NS).get('isExecutable')=='false'
        for e in root.iter():
            for attr in ('sourceRef','targetRef','attachedToRef','processRef','bpmnElement','default','errorRef'):
                if e.get(attr): assert e.get(attr) in elements, f'{m.name}: missing {attr}'
            if e.tag in ('{%s}incoming'%NS['bpmn'],'{%s}outgoing'%NS['bpmn'],'{%s}flowNodeRef'%NS['bpmn']):
                assert e.text in elements, f'{m.name}: missing node/flow'
        lane_ids=[e.text for e in root.findall('.//bpmn:flowNodeRef',NS)]
        assert sorted(lane_ids)==sorted(m.nodes), f'{m.name}: lane membership'
        start=[n['id'] for n in m.nodes.values() if n['kind']=='startEvent']
        assert len(start)==1
        reachable=set(start)
        while True:
            before=len(reachable)
            for f in m.flows:
                if f['source'] in reachable: reachable.add(f['target'])
            for n in m.nodes.values():
                if n.get('attachedToRef') in reachable: reachable.add(n['id'])
            if len(reachable)==before: break
        assert reachable==set(m.nodes), f'{m.name}: unreachable {set(m.nodes)-reachable}'
        terminating={n['id'] for n in m.nodes.values() if n['kind']=='endEvent'}
        while True:
            before=len(terminating)
            for f in m.flows:
                if f['target'] in terminating: terminating.add(f['source'])
            if len(terminating)==before: break
        assert terminating==set(m.nodes), f'{m.name}: no exit'
        for n in m.nodes.values():
            incoming=[f for f in m.flows if f['target']==n['id']]
            outgoing=[f for f in m.flows if f['source']==n['id']]
            assert incoming or n['kind'] in ('startEvent','boundaryEvent')
            assert outgoing or n['kind']=='endEvent'
            if n['kind']=='exclusiveGateway' and len(outgoing)>1:
                assert sum(f['default'] for f in outgoing)==1
                assert all(f['label'] for f in outgoing)
            lane=next(l for l in m.lanes if l[0]==n['lane'])
            assert 110<n['x'] and n['x']+n['w']<1645 and lane[2]<=n['y'] and n['y']+n['h']<=lane[2]+lane[3]
        shape_ids={e.get('bpmnElement') for e in root.findall('.//bpmndi:BPMNShape',NS)}
        assert set(m.nodes).issubset(shape_ids)
        edge_ids={e.get('bpmnElement') for e in root.findall('.//bpmndi:BPMNEdge',NS)}
        assert edge_ids=={f['id'] for f in m.flows+m.messages}
        for f in m.messages:
            assert any(f[side] in {p[0] for p in m.external} for side in ('source','target'))
        for f in m.flows:
            for a,b in zip(f['points'],f['points'][1:]):
                assert a[0]==b[0] or a[1]==b[1], f'{m.name}: non-orthogonal edge'
                for node in m.nodes.values():
                    if node['id'] in (f['source'],f['target']) or node.get('attachedToRef') in (f['source'],f['target']):
                        continue
                    x,y,w,h=node['x'],node['y'],node['w'],node['h']
                    hit=(a[0]==b[0] and x+2<a[0]<x+w-2 and max(min(a[1],b[1]),y+2)<min(max(a[1],b[1]),y+h-2)) or (a[1]==b[1] and y+2<a[1]<y+h-2 and max(min(a[0],b[0]),x+2)<min(max(a[0],b[0]),x+w-2))
                    assert not hit, f'{m.name}: edge through {node["id"]}'
            assert all(110<=x<=1645 and m.lanes[0][2]<=y<=m.lanes[-1][2]+m.lanes[-1][3] for x,y in f['points'])
        for suffix in ('.svg','.png'): assert (OUT/(m.name+suffix)).stat().st_size>1000
        ET.parse(OUT/(m.name+'.svg'))
        with Image.open(OUT/(m.name+'.png')) as im: assert im.size==(3400,m.height*2)
        if schema_dir:
            from lxml import etree
            schema=etree.XMLSchema(etree.parse(str(Path(schema_dir)/'BPMN20.xsd')))
            schema.assertValid(etree.parse(str(OUT/(m.name+'.bpmn'))))
        print(m.name+': XML, lanes, branches, reachability, exits, DI and images OK')
    source=(HERE/'README.md').read_text(encoding='utf-8')
    for prefix,count in [('FR',14),('NFR',10),('BR',12)]:
        for i in range(1,count+1): assert f'| {prefix}{i:02}' in source
    report=(ROOT/'docs/report/BaoCao.md').read_text(encoding='utf-8')
    assert 'Nhóm 2B' in report and 'Nhóm 5B' not in report
    assert '### 3.1' in report and '### 4.5' in report
    assert report.count('## Chương ')==7
    for image in re.findall(r'!\[[^\]]*\]\(([^)]+)\)',report):
        assert (ROOT/'docs/report'/image).is_file(), image
    from docx import Document
    doc=Document(ROOT/'docs/report/BaoCao.docx')
    text='\n'.join(p.text for p in doc.paragraphs)
    assert '3.1 Phạm vi' in text and '4.5 Kiểm tra mô hình' in text
    assert len(doc.inline_shapes)>=5 and len(doc.tables)>=8
    assert sum(p.text.startswith('Chương ') for p in doc.paragraphs)==7
    print('Report: seven chapters, requirement IDs, tables and Week 04 figures OK')


def sync_report():
    path=ROOT/'docs/report/BaoCao.md'
    old=path.read_text(encoding='utf-8')
    content=(HERE/'README.md').read_text(encoding='utf-8')
    chapters=content[content.index('## Chương 3'):content.index('## Sử dụng bộ hồ sơ')].strip()
    chapters=chapters.replace('](diagrams/','](../week04/diagrams/')
    # Keep the later Week 05 addition when rebuilding Week 04 figures.
    if '### 3.10 ' in old:
        week05_ch3=old[old.index('### 3.10 '):old.index('## Chương 4')].strip()
        chapters=chapters.replace('## Chương 4',week05_ch3+'\n\n## Chương 4',1)
    if '### 4.6 ' in old:
        chapters+='\n\n'+old[old.index('### 4.6 '):old.index('## Chương 5')].strip()
    updated=old[:old.index('## Chương 3')]+chapters+'\n\n'+old[old.index('## Chương 5'):]
    updated=updated.replace('Báo cáo tiến độ ngày 8 tháng 10 năm 2026',
                            'Báo cáo tiến độ và hồ sơ tuần 04 cập nhật ngày 15 tháng 10 năm 2026')
    updated=updated.replace('Nhóm 5B','Nhóm 2B')
    updated=updated.replace('BPMN chuẩn nếu biểu mẫu yêu cầu và kết quả đo tải',
                            'duyệt nghiệp vụ cho bộ BPMN tuần 04 và kết quả đo tải')
    path.write_text(updated,encoding='utf-8')
    make_docx(updated)


def make_docx(markdown):
    from docx import Document
    from docx.shared import Mm,Pt,RGBColor
    from docx.enum.section import WD_SECTION,WD_ORIENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    path=ROOT/'docs/report/BaoCao.docx'
    d=Document(path)
    body=d._element.body
    for child in list(body):
        if child.tag!=qn('w:sectPr'): body.remove(child)
    for id,rel in list(d.part.rels.items()):
        if rel.reltype.endswith('/image'): d.part.drop_rel(id)
    # Preserve original Times New Roman report typography and A4 margins.
    for name in ('Normal','Title','Heading 1','Heading 2','Heading 3','Caption'):
        style=d.styles[name]; style.font.name='Times New Roman'; style.font.color.rgb=RGBColor(0,0,0)
        if name.startswith('Heading'): style.paragraph_format.keep_with_next=True
    for sec in d.sections:
        sec.page_width=Mm(210); sec.page_height=Mm(297); sec.orientation=WD_ORIENT.PORTRAIT
        sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Mm(20)
    footer=d.sections[0].footer.paragraphs[0]; footer.clear(); footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)
    def plain(text):
        return re.sub(r'\[([^\]]+)\]\([^)]+\)',r'\1',text).replace('**','').replace('`','')
    def textparagraph(text,style=None):
        p=d.add_paragraph(plain(text),style)
        p.paragraph_format.widow_control=True
        return p
    def landscape():
        s=d.add_section(WD_SECTION.NEW_PAGE); s.orientation=WD_ORIENT.LANDSCAPE
        s.page_width=Mm(297); s.page_height=Mm(210)
        s.top_margin=s.bottom_margin=Mm(15); s.left_margin=s.right_margin=Mm(15)
        return s
    def portrait():
        s=d.add_section(WD_SECTION.NEW_PAGE); s.orientation=WD_ORIENT.PORTRAIT
        s.page_width=Mm(210); s.page_height=Mm(297)
        s.top_margin=s.bottom_margin=s.left_margin=s.right_margin=Mm(20)
    lines=markdown.splitlines(); i=0; in_landscape=False
    while i<len(lines):
        line=lines[i].strip(); i+=1
        if not line: continue
        if in_landscape and not line.startswith('!['):
            portrait(); in_landscape=False
        if line.startswith('!['):
            match=re.fullmatch(r'!\[([^\]]+)\]\(([^)]+)\)',line)
            assert match
            if in_landscape: d.add_page_break()
            else: landscape(); in_landscape=True
            from PIL import Image
            image=ROOT/'docs/report'/match[2]
            with Image.open(image) as im: ratio=im.height/im.width
            width=min(267,(160 if 'week05' in image.parts else 151)/ratio)
            p=d.add_paragraph(); p.paragraph_format.space_after=Pt(0); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(image),width=Mm(width))
            caption_text=match[1] if match[1].startswith('Hình ') else 'Hình '+('4.1' if image.stem=='BFD' else '4.'+str(int(image.stem[-2:])+1))+' '+match[1]
            caption=textparagraph(caption_text,'Caption')
            caption.alignment=WD_ALIGN_PARAGRAPH.CENTER; caption.paragraph_format.space_after=Pt(0)
        elif line.startswith('|'):
            rows=[line]
            while i<len(lines) and lines[i].strip().startswith('|'):
                rows.append(lines[i].strip()); i+=1
            values=[[plain(c.strip()) for c in r.strip('|').split('|')] for r in rows]
            values=[r for r in values if not all(re.fullmatch(r'[:\- ]+',c) for c in r)]
            columns=len(values[0]); assert all(len(r)==columns for r in values)
            table=d.add_table(rows=0,cols=columns); table.autofit=False
            ratios={2:[.14,.86],3:[.18,.39,.43],4:[.15,.23,.30,.32]}[columns]
            if values[0]==['Trường','Đặc tả']:
                ratios=[.23,.77]
            if columns==4 and values[0][2]=='Mục tiêu và cách kiểm chứng':
                ratios=[.12,.22,.41,.25]
            for c,fraction in zip(table.columns,ratios): c.width=Mm(170*fraction)
            for row_index,values_row in enumerate(values):
                row=table.add_row()
                trpr=row._tr.get_or_add_trPr()
                trpr.append(OxmlElement('w:cantSplit'))
                if row_index==0: trpr.append(OxmlElement('w:tblHeader'))
                for j,(cell,value) in enumerate(zip(row.cells,values_row)):
                    cell.width=Mm(170*ratios[j]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    tcpr=cell._tc.get_or_add_tcPr(); margin=OxmlElement('w:tcMar')
                    for side in ('top','left','bottom','right'):
                        element=OxmlElement('w:'+side); element.set(qn('w:w'),'80'); element.set(qn('w:type'),'dxa'); margin.append(element)
                    tcpr.append(margin)
                    shading=OxmlElement('w:shd'); shading.set(qn('w:fill'),'E8EDF2' if row_index==0 else 'FFFFFF'); tcpr.append(shading)
                    border=OxmlElement('w:tcBorders')
                    for side in ('top','left','bottom','right'):
                        element=OxmlElement('w:'+side); element.set(qn('w:val'),'single'); element.set(qn('w:sz'),'4'); element.set(qn('w:color'),'D9D9D9'); border.append(element)
                    tcpr.append(border)
                    p=cell.paragraphs[0]; p.paragraph_format.space_before=Pt(2); p.paragraph_format.space_after=Pt(2)
                    p.paragraph_format.line_spacing=1.05
                    p.paragraph_format.keep_with_next=(row_index<len(values)-1) if values[0]==['Trường','Đặc tả'] else row_index==0
                    if j==0 and len(value)<12: p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                    run=p.add_run(value); run.font.size=Pt(10.5); run.bold=row_index==0
            d.add_paragraph().paragraph_format.space_after=Pt(0)
        elif line.startswith('# '): textparagraph(line[2:],'Title')
        elif line.startswith('## '): textparagraph(line[3:],'Heading 1')
        elif line.startswith('### '): textparagraph(line[4:],'Heading 2')
        elif line.startswith('#### '): textparagraph(line[5:],'Heading 3')
        elif re.match(r'\[S\d+\]',line):
            p=textparagraph(line); p.paragraph_format.line_spacing=1.0; p.paragraph_format.space_after=Pt(6)
            for run in p.runs: run.font.size=Pt(11)
        else: textparagraph(line)
    d.core_properties.title='Nền tảng trợ lý AI hỗ trợ khách hàng đa kênh dựa trên kiến trúc RAG'
    d.core_properties.subject='Chương 3 và 4 Hồ sơ tuần 04 và 05 Nhóm 2B' if '### 3.10 ' in markdown else 'Chương 3 và 4 Hồ sơ tuần 04 Nhóm 2B'
    d.core_properties.author='Nhóm 2B'
    temp=path.with_name('BaoCao.week04.tmp.docx')
    d.save(temp)
    # Atomic replacement avoids partial overwrite if report is locked in Word.
    try: temp.replace(path)
    except OSError:
        raise RuntimeError(f'Cannot replace locked report; completed copy retained at {temp}') from None


if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--check',action='store_true')
    parser.add_argument('--schema-dir',type=Path)
    args=parser.parse_args()
    if not args.check:
        OUT.mkdir(parents=True,exist_ok=True)
        draw_bfd()
        for model in models(): write_bpmn(model); draw_model(model)
        sync_report()
    check(args.schema_dir)
